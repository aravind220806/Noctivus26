from copy import deepcopy

import pytest
import pytest_asyncio

from app.services import scheduler_service as scheduler
from app.services import event_service
from app.db import sqlite_db as db_module


@pytest_asyncio.fixture(params=['memory', 'sqlite'])
async def store(request, monkeypatch, tmp_path):
    db = db_module._SQLiteDB()
    monkeypatch.setattr(db_module, 'DB_PATH', tmp_path / 'scheduler.db')
    monkeypatch.setattr(scheduler, 'sqlite_db', db)
    monkeypatch.setattr(event_service, 'sqlite_db', db)
    monkeypatch.setattr(scheduler, 'memory_event_slots', [])
    monkeypatch.setattr(event_service, 'memory_events', [])
    monkeypatch.setattr(scheduler, '_last_assignment_summary', None)
    if request.param == 'sqlite':
        await db.init()
    rows = {}

    async def load(filters=None):
        return [deepcopy(row) for row in rows.values()
                if not filters or all(row.get('paymentStatus' if key == 'status' else key) == value for key, value in filters.items())]

    async def update(member_id, changes):
        rows[member_id].update(deepcopy(changes))
        return deepcopy(rows[member_id])

    monkeypatch.setattr(scheduler, 'load_registrations', load)
    monkeypatch.setattr(scheduler, 'update_registration', update)
    return rows


def registration(member_id, events, assigned=()):
    return {'registrationId': member_id, 'paymentStatus': 'confirmed', 'event_ids': events,
            'eventRegistrations': [{'eventId': eid} for eid in events], 'assigned_slots': list(assigned)}


def slot(slot_id, event='bug-hunt', start='10:00', end='13:00', members=(), capacity=30):
    return {'id': slot_id, 'event_id': event, 'date': '2026-09-26', 'window': 'morning',
            'start_time': start, 'end_time': end, 'capacity': capacity, 'assigned_member_ids': list(members)}


@pytest.mark.asyncio
async def test_policy_combines_single_event_and_splits_other_event(store):
    for index in range(65):
        mid = f'm{index}'
        store[mid] = registration(mid, ['bug-hunt', 'tune-trap'])
    result = await scheduler.repair_event_schedule()
    slots = await scheduler.load_all_slots()
    bug_slots = [slot for slot in slots if slot['event_id'] == 'bug-hunt']
    tune_slots = [slot for slot in slots if slot['event_id'] == 'tune-trap']
    assert len(bug_slots) == 1
    assert len(tune_slots) == 2
    assert len(bug_slots[0]['assigned_member_ids']) == 65
    assert result['successfully_assigned'] == 65
    assert result['unassigned_conflicts'] == []
    assert result['unassigned_full'] == []
    assert (await event_service.get_event('bug-hunt'))['slot_count'] == 1
    assert (await event_service.get_event('tune-trap'))['slot_count'] == 2


@pytest.mark.asyncio
async def test_delete_reassigns_partial_registration_and_preserves_other_event(store):
    store['member'] = registration('member', ['bug-hunt', 'tune-trap'], ['deleted', 'other'])
    await scheduler.save_slot(slot('deleted', members=['member']))
    await scheduler.save_slot(slot('remaining'))
    await scheduler.save_slot(slot('other', 'tune-trap', '13:00', '16:00', ['member']))
    assert await scheduler.delete_slot('deleted')
    assert set(store['member']['assigned_slots']) == {'remaining', 'other'}
    assert store['member']['eventRegistrations'][0]['batchTime'] == '10:00'
    remaining = await scheduler.get_slot('remaining')
    assert remaining['assigned_member_ids'] == ['member']
    await scheduler.assignMembersToSlots(auto_generate=False)
    assert (await scheduler.get_slot('remaining'))['assigned_member_ids'] == ['member']


@pytest.mark.asyncio
async def test_delete_combines_members_and_expands_remaining_capacity(store):
    for index in range(40):
        mid = f'm{index}'
        store[mid] = registration(mid, ['bug-hunt'], ['deleted' if index < 20 else 'remaining'])
    await scheduler.save_slot(slot('deleted', members=[f'm{i}' for i in range(20)]))
    await scheduler.save_slot(slot('remaining', members=[f'm{i}' for i in range(20, 40)]))
    await scheduler.delete_slot('deleted')
    slots = await scheduler.load_all_slots()
    assert len(slots) == 1
    assert slots[0]['capacity'] >= 40
    assert len(slots[0]['assigned_member_ids']) == 40


@pytest.mark.asyncio
async def test_delete_last_slot_does_not_recreate_on_dashboard_load(store):
    store['member'] = registration('member', ['bug-hunt'], ['deleted'])
    await scheduler.save_slot(slot('deleted', members=['member']))
    await scheduler.delete_slot('deleted')
    await scheduler.get_scheduler_dashboard_data()
    assert await scheduler.load_all_slots() == []
    assert store['member']['assigned_slots'] == []
    assert scheduler._last_assignment_summary['unassigned_full'] == ['member']


@pytest.mark.asyncio
async def test_conflict_is_reported_without_creating_extra_slots(store):
    store['member'] = registration('member', ['bug-hunt', 'tune-trap'], ['deleted', 'other'])
    await scheduler.save_slot(slot('deleted', members=['member']))
    await scheduler.save_slot(slot('remaining'))
    await scheduler.save_slot(slot('other', 'tune-trap', members=['member']))
    await scheduler.delete_slot('deleted')
    assert store['member']['assigned_slots'] == ['other']
    assert scheduler._last_assignment_summary['unassigned_conflicts'] == ['member']
    assert len(await scheduler.load_all_slots()) == 2


@pytest.mark.asyncio
async def test_move_other_assignment_when_needed_for_valid_pair(store):
    store['member'] = registration('member', ['bug-hunt', 'tune-trap'], ['deleted', 'other'])
    for item in [slot('deleted', members=['member']), slot('remaining'),
                 slot('other', 'tune-trap', members=['member']),
                 slot('later', 'tune-trap', '13:00', '16:00')]:
        await scheduler.save_slot(item)
    await scheduler.delete_slot('deleted')
    assert set(store['member']['assigned_slots']) == {'remaining', 'later'}
    assert (await scheduler.get_slot('other'))['assigned_member_ids'] == []


@pytest.mark.asyncio
async def test_full_custom_slots_are_reported(store):
    store['existing'] = registration('existing', ['bug-hunt'], ['only'])
    store['waiting'] = registration('waiting', ['bug-hunt'])
    await scheduler.save_slot(slot('only', members=['existing'], capacity=1))
    result = await scheduler.assignMembersToSlots(auto_generate=False)
    assert result['unassigned_full'] == ['waiting']
    assert len(await scheduler.load_all_slots()) == 1


@pytest.mark.parametrize('count', [0, 3, True, '2', None])
@pytest.mark.asyncio
async def test_invalid_slot_count_rejected_without_mutation(store, count):
    with pytest.raises(ValueError):
        await scheduler.configure_event_slots('bug-hunt', count, 'admin')
    assert await scheduler.load_all_slots() == []


@pytest.mark.parametrize('event_id', ['bug-hunt', 'prompt-heist', 'ignite', 'ctf', 'tune-trap'])
@pytest.mark.parametrize('count', [1, 2])
def test_generation_respects_choice_regardless_of_registration_count(event_id, count):
    slots = scheduler.generateSlotsForEvent({'id': event_id, 'slot_count': count, 'duration_minutes': 180}, 200)
    assert len(slots) == scheduler.required_slot_count(event_id)
    assert slots[0]['start_time'] == '10:00'
    assert sum(s['capacity'] for s in slots) >= 200
    if len(slots) == 2:
        assert not scheduler.slotsConflict(*slots)


@pytest.mark.asyncio
async def test_splitting_remaining_afternoon_slot_uses_distinct_ids(store):
    store['member'] = registration('member', ['tune-trap'])
    await scheduler.configure_event_slots('tune-trap', 2, 'admin')
    await scheduler.delete_slot('slot_tune-trap_morning_1')
    await scheduler.configure_event_slots('tune-trap', 2, 'admin')
    slots = await scheduler.load_all_slots()
    assert len(slots) == 2
    assert len({slot['id'] for slot in slots}) == 2
    assert not scheduler.slotsConflict(*slots)


@pytest.mark.asyncio
async def test_new_members_fit_configured_single_slot_without_expansion(store):
    await scheduler.configure_event_slots('bug-hunt', 1, 'admin')
    for index in range(40):
        mid = f'm{index}'
        store[mid] = registration(mid, ['bug-hunt'])
    summary = await scheduler.assignMembersToSlots(auto_generate=False)
    assert summary['successfully_assigned'] == 40
    assert len(await scheduler.load_all_slots()) == 1


def assert_valid_assignments(rows, slots):
    by_id = {slot['id']: slot for slot in slots}
    for mid, reg in rows.items():
        assigned = [by_id[sid] for sid in reg['assigned_slots']]
        assert {s['event_id'] for s in assigned} == set(scheduler._get_event_ids(reg))
        for index, left in enumerate(assigned):
            assert mid in left['assigned_member_ids']
            for right in assigned[index + 1:]:
                assert not scheduler.slotsConflict(left, right)
    for slot in slots:
        assert len(slot['assigned_member_ids']) == len(set(slot['assigned_member_ids']))
        assert len(slot['assigned_member_ids']) <= slot['capacity']


@pytest.mark.asyncio
async def test_auto_upgrade_repairs_112_conflicting_members(store):
    from app.events import EVENT_CATALOG
    for event in EVENT_CATALOG:
        await scheduler.save_slot(slot(f"old-{event['id']}", event['id']))
    for index in range(112):
        tech = ['bug-hunt', 'prompt-heist', 'ignite', 'secure-x-vibecode'][index % 4]
        other = ['mystery-hunt', 'tune-trap', 'ipl-bidverse'][index % 3]
        mid = f'm{index}'
        # Include fully assigned but overlapping members as well as unassigned.
        assigned = [f'old-{tech}', f'old-{other}'] if index % 2 else []
        store[mid] = registration(mid, [tech, other], assigned)
    dashboard = await scheduler.get_scheduler_dashboard_data()
    summary = dashboard['last_assignment_summary']
    assert summary['successfully_assigned'] == 112
    assert summary['unassigned_conflicts'] == []
    assert summary['unassigned_full'] == []
    for event in dashboard['events']:
        assert event['slots_count'] == (1 if event['id'] in {'ctf', 'bug-hunt', 'prompt-heist', 'ignite', 'playground-of-hackers'} else 2)
    slots = await scheduler.load_all_slots()
    assert_valid_assignments(store, slots)
    await scheduler.get_scheduler_dashboard_data()
    assert await scheduler.load_all_slots() == slots


@pytest.mark.asyncio
async def test_staggers_single_sessions_with_shared_members(store):
    for index, pair in enumerate([['bug-hunt', 'prompt-heist'], ['prompt-heist', 'ignite'], ['ignite', 'bug-hunt']]):
        store[str(index)] = registration(str(index), pair)
    summary = await scheduler.repair_event_schedule()
    assert summary['unassigned_conflicts'] == []
    assert_valid_assignments(store, await scheduler.load_all_slots())


@pytest.mark.asyncio
async def test_workshop_and_other_event_keep_duration_and_have_valid_pair(store):
    store['member'] = registration('member', ['playground-of-hackers', 'bug-hunt'])
    summary = await scheduler.repair_event_schedule()
    assert_valid_assignments(store, await scheduler.load_all_slots())
    workshops = [s for s in await scheduler.load_all_slots() if s['event_id'] == 'playground-of-hackers']
    assert len(workshops) == 1
    assert workshops[0]['start_time'] == '10:00'
    assert workshops[0]['end_time'] == '16:00'
    assert not any(s['event_id'] == 'playground-of-hackers' for s in summary['late_sessions'])
    event = await event_service.get_event('playground-of-hackers')
    assert event['duration_minutes'] == 360


@pytest.mark.asyncio
async def test_standalone_members_are_balanced_between_two_slots(store):
    for index in range(65):
        store[str(index)] = registration(str(index), ['tune-trap'])
    await scheduler.repair_event_schedule()
    slots = [s for s in await scheduler.load_all_slots() if s['event_id'] == 'tune-trap']
    assert sorted(len(s['assigned_member_ids']) for s in slots) == [32, 33]


@pytest.mark.asyncio
async def test_invalid_plan_leaves_existing_schedule_and_members_unchanged(store, monkeypatch):
    from unittest.mock import AsyncMock
    store['member'] = registration('member', ['tune-trap'], ['old'])
    await scheduler.save_slot(slot('old', 'tune-trap', members=['member']))
    before = deepcopy(store)
    slots_before = await scheduler.load_all_slots()
    monkeypatch.setattr(scheduler, 'list_events', AsyncMock(return_value=[{'id': 'tune-trap', 'duration_minutes': 1000}]))
    with pytest.raises(ValueError):
        await scheduler.repair_event_schedule()
    assert store == before
    assert await scheduler.load_all_slots() == slots_before


@pytest.mark.asyncio
async def test_rechecks_complete_but_overlapping_assignments(store):
    store['member'] = registration('member', ['bug-hunt', 'tune-trap'], ['bug', 'morning'])
    for s in [slot('bug'), slot('morning', 'tune-trap'), slot('afternoon', 'tune-trap', '13:00', '16:00')]:
        await scheduler.save_slot(s)
    summary = await scheduler.assignMembersToSlots(auto_generate=False)
    assert summary['successfully_assigned'] == 1
    assert_valid_assignments(store, await scheduler.load_all_slots())


@pytest.mark.asyncio
async def test_dashboard_discards_stale_conflict_summary(store, monkeypatch):
    store['member'] = registration('member', ['bug-hunt', 'tune-trap'])
    await scheduler.repair_event_schedule()
    monkeypatch.setattr(scheduler, '_last_assignment_summary', {
        'total_processed': 112, 'successfully_assigned': 0, 'unassigned_conflicts': ['member'], 'unassigned_full': []
    })
    dashboard = await scheduler.get_scheduler_dashboard_data()
    assert dashboard['last_assignment_summary']['successfully_assigned'] == 1
    assert dashboard['last_assignment_summary']['unassigned_conflicts'] == []


@pytest.mark.asyncio
async def test_cannot_add_extra_slots_above_event_policy(store):
    await scheduler.repair_event_schedule()
    with pytest.raises(ValueError, match='limited to 1'):
        await scheduler.create_custom_slot({'event_id': 'bug-hunt'})
    with pytest.raises(ValueError, match='limited to 2'):
        await scheduler.create_custom_slot({'event_id': 'tune-trap'})


@pytest.mark.asyncio
async def test_old_two_session_workshop_is_migrated_to_fixed_window(store):
    store['member'] = registration('member', ['playground-of-hackers'], ['old-workshop'])
    old = slot('old-workshop', 'playground-of-hackers', '15:00', '20:00', ['member'])
    old['schedule_policy'] = 'fixed-counts-conflict-repair-v1'
    await scheduler.save_slot(old)
    dashboard = await scheduler.get_scheduler_dashboard_data()
    workshop = next(e for e in dashboard['events'] if e['id'] == 'playground-of-hackers')
    assert workshop['duration_minutes'] == 360
    assert workshop['slots_count'] == 1
    assert workshop['slots'][0]['start_time'] == '10:00'
    assert workshop['slots'][0]['end_time'] == '16:00'
    assert_valid_assignments(store, await scheduler.load_all_slots())


@pytest.mark.asyncio
async def test_workshop_window_cannot_be_changed_by_slot_edit(store):
    await scheduler.repair_event_schedule()
    workshop = next(s for s in await scheduler.load_all_slots() if s['event_id'] == 'playground-of-hackers')
    with pytest.raises(ValueError, match='10:00 AM to 4:00 PM'):
        await scheduler.update_slot(workshop['id'], {'start_time': '11:00'})
    assert (await scheduler.get_slot(workshop['id']))['start_time'] == '10:00'


@pytest.mark.asyncio
async def test_supplied_csv_event_distribution_assigns_all_205_members(store):
    import json
    from pathlib import Path
    fixture = json.loads((Path(__file__).parent / 'fixtures' / 'scheduler_registration_counts.json').read_text())
    for combination in fixture['combinations']:
        for _ in range(combination['count']):
            member_id = f'ANONYMOUS-{len(store) + 1:04d}'
            store[member_id] = registration(member_id, combination['event_ids'])
    summary = await scheduler.repair_event_schedule()
    slots = await scheduler.load_all_slots()
    assert len(store) == fixture['confirmed_count'] == 205
    assert summary['successfully_assigned'] == 205
    assert summary['unassigned_conflicts'] == []
    assert summary['unassigned_full'] == []
    assert summary['late_sessions'] == []
    assert sum(len(reg['assigned_slots']) for reg in store.values()) == 306
    assert len(slots) == 13
    assert_valid_assignments(store, slots)
    for slot in slots:
        assert slot['start_time'] >= '10:00'
        assert slot['end_time'] <= '16:00'
    workshop = next(s for s in slots if s['event_id'] == 'playground-of-hackers')
    assert (workshop['start_time'], workshop['end_time']) == ('10:00', '16:00')
    assert len(workshop['assigned_member_ids']) == 12


@pytest.mark.asyncio
async def test_excel_exports_saved_reason_for_each_event(store):
    from io import BytesIO
    from openpyxl import load_workbook
    from app.services.export_service import export_scheduler_to_excel
    store['member'] = registration('member', ['ignite', 'tune-trap'])
    await scheduler.repair_event_schedule()
    slots = await scheduler.load_all_slots()
    reg = store['member']
    assert set(reg['slot_assignment_reasons']) == set(reg['assigned_slots'])
    workbook = load_workbook(BytesIO(export_scheduler_to_excel(await event_service.list_events(), slots, [reg])))
    sheet = workbook['Event Member Schedule']
    headers = [cell.value for cell in sheet[1]]
    reason_column = headers.index('Why This Batch')
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 2
    assert all(row[reason_column] for row in rows)
    tune = next(row for row in rows if row[1] == 'Tune Trap')
    assert 'overlaps IGNITE (10:00–13:00)' in tune[reason_column]
    assert workbook['Member Allocations'].cell(1, 11).value == 'Batch Assignment Reasons'
    assert 'Tune Trap:' in workbook['Member Allocations'].cell(2, 11).value
