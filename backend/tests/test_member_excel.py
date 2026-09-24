from io import BytesIO
from openpyxl import load_workbook
from app.services.export_service import export_members_to_excel


def test_member_events_are_separate_and_safe():
    registrations = [
        {'registrationId': 'NOC26-1', 'participant': {'name': '=BAD()', 'phone': '0123456789'}, 'paymentStatus': 'confirmed', 'eventRegistrations': [{'eventName': 'Cyber Challenge'}, {'eventId': 'quiz'}]},
        {'registrationId': 'NOC26-2', 'participant': {'name': 'Single Event'}, 'eventRegistrations': [{'eventName': 'Quiz'}]},
    ]
    sheet = load_workbook(BytesIO(export_members_to_excel(registrations, [{'id': 'quiz', 'name': 'Technical Quiz'}]))).active
    rows = list(sheet.values)
    assert rows[0][5:7] == ('Event 1', 'Event 2')
    assert rows[1][5:7] == ('Cyber Challenge', 'Technical Quiz')
    assert rows[1][1] == "'=BAD()"
    assert sheet['B2'].data_type == 's'
    assert rows[1][3] == '0123456789'
    assert rows[2][5:7] == ('Quiz', None)
    assert sheet.freeze_panes == 'A2'
    assert sheet.auto_filter.ref == 'A1:H3'


def test_empty_export_and_extra_events():
    empty = load_workbook(BytesIO(export_members_to_excel([], []))).active
    assert list(empty.values)[0][5:7] == ('Event 1', 'Event 2')
    sheet = load_workbook(BytesIO(export_members_to_excel([{'eventRegistrations': [{'eventName': name} for name in ['One', 'Two', 'Three']]}], []))).active
    assert list(sheet.values)[1][5:8] == ('One', 'Two', 'Three')


def test_filtered_excel_route_keeps_both_event_names(monkeypatch):
    from unittest.mock import AsyncMock
    from fastapi.testclient import TestClient
    from app.main import app
    from app.routes import admin_routes
    from app.middleware.admin_auth import sign_admin_token
    import app.middleware.admin_auth as auth
    loader = AsyncMock(return_value=[{'registrationId': 'NOC26-1', 'participant': {'name': 'Member'}, 'eventRegistrations': [{'eventName': 'First Event'}, {'eventName': 'Second Event'}]}])
    monkeypatch.setattr(admin_routes, 'load_registrations', loader)
    monkeypatch.setattr(admin_routes, 'list_events', AsyncMock(return_value=[]))
    monkeypatch.setattr(admin_routes, 'record_admin_action', AsyncMock())
    monkeypatch.setattr(auth, 'resolve_admin_access', AsyncMock(return_value={'tabs': ['Export'], 'owner': False}))
    monkeypatch.setattr(auth, 'session_exists', AsyncMock(return_value=True))
    client = TestClient(app)
    path = '/api/admin/export/members-excel?eventId=first&status=confirmed'
    assert client.get(path).status_code == 401
    token, _ = sign_admin_token({'email': 'export@example.com', 'name': 'Export Admin', 'tabs': ['Export'], 'owner': False})
    client.cookies.set('noctivus_admin_session', token)
    response = client.get(path)
    assert response.status_code == 200
    loader.assert_awaited_once_with({'eventId': 'first', 'status': 'confirmed'})
    sheet = load_workbook(BytesIO(response.content)).active
    assert list(sheet.values)[1][5:7] == ('First Event', 'Second Event')
    assert '.xlsx' in response.headers['content-disposition']
